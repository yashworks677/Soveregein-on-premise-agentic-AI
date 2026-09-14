import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from pathlib import Path
from datetime import datetime
from typing import Dict, Any, List

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docx_approval_report(data: Dict[str, Any], output_path: Path) -> str:
    """
    Generates a formal, enterprise-grade DOCX technical approval memorandum.
    """
    doc = docx.Document()

    # Configure Margins (0.8 inch)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header / Banner
    header_para = doc.add_paragraph()
    header_run = header_para.add_run("SOVEREIGN INDUSTRIAL & DEFENCE ENCLAVE")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(9)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor(71, 85, 105)
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Title
    title_para = doc.add_paragraph()
    title_run = title_para.add_run(f"TECHNICAL EVALUATION & APPROVAL MEMORANDUM")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42)

    sub_para = doc.add_paragraph()
    sub_run = sub_para.add_run(f"Task: {data.get('task_type', 'Document Analysis')} | Requirement: {data.get('requirement', '')[:90]}")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(10)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Enclave Clearance Level", "CONFIDENTIAL / INTERNAL OPERATIONAL USE ONLY"),
        ("Processing Infrastructure", "Air-Gapped Sovereign Local Host (Zero External APIs)"),
        ("Source Documents Analyzed", ", ".join(data.get("input_files", [])) or "Operational Telemetry"),
        ("Verification Timestamp", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"))
    ]
    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.text = label
        cell_lbl.paragraphs[0].runs[0].font.bold = True
        cell_lbl.paragraphs[0].runs[0].font.size = Pt(9)
        cell_lbl.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 41, 59)
        set_cell_background(cell_lbl, "F1F5F9")
        
        cell_val.text = val
        cell_val.paragraphs[0].runs[0].font.size = Pt(9)
        cell_val.paragraphs[0].runs[0].font.color.rgb = RGBColor(51, 65, 85)
        set_cell_background(cell_val, "F8FAFC")

    doc.add_paragraph()

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Compliance Verdict", level=1)
    h1.style.font.color.rgb = RGBColor(15, 23, 42)
    p_summary = doc.add_paragraph(data.get("executive_summary", "No summary provided."))
    p_summary.style.font.size = Pt(10.5)

    # Section 2: Key Technical Findings
    h2 = doc.add_heading("2. Verified Technical Findings", level=1)
    h2.style.font.color.rgb = RGBColor(15, 23, 42)
    for finding in data.get("key_findings", []):
        p_f = doc.add_paragraph(style='List Bullet')
        r_f = p_f.add_run(finding)
        r_f.font.size = Pt(10)

    # Section 3: Critical Issues & Deviations
    critical_issues = data.get("critical_issues", [])
    if critical_issues:
        h3 = doc.add_heading("3. Critical Safety & Integrity Deviations", level=1)
        h3.style.font.color.rgb = RGBColor(15, 23, 42)

        issue_table = doc.add_table(rows=1, cols=4)
        issue_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        headers = ["Severity", "Component / Tag", "Observed Deviation", "Safety Risk"]
        for idx, text in enumerate(headers):
            cell = issue_table.rows[0].cells[idx]
            cell.text = text
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(cell, "1E293B")

        for issue in critical_issues:
            row = issue_table.add_row()
            sev = issue.get("severity", "WARNING")
            row.cells[0].text = sev
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[0].paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(row.cells[0], "FEE2E2" if sev == "CRITICAL" else "FEF3C7")

            row.cells[1].text = issue.get("component", issue.get("column", "System Wide"))
            row.cells[1].paragraphs[0].runs[0].font.size = Pt(9)

            row.cells[2].text = issue.get("deviation", issue.get("message", ""))
            row.cells[2].paragraphs[0].runs[0].font.size = Pt(9)

            row.cells[3].text = issue.get("risk", issue.get("expected_range", "Requires Engineering Action"))
            row.cells[3].paragraphs[0].runs[0].font.size = Pt(9)

    # Section 4: Sources Used / Standards Referenced
    sources = data.get("sources_used", [])
    if sources:
        doc.add_paragraph()
        h4 = doc.add_heading("4. Internal Knowledge Base References (Qdrant RAG)", level=1)
        h4.style.font.color.rgb = RGBColor(15, 23, 42)
        for src in sources:
            p_s = doc.add_paragraph(style='List Bullet')
            r_s = p_s.add_run(f"{src['document_name']} ({src.get('section', 'General')}) - Relevance: {int(src['relevance_score']*100)}%\n")
            r_s.font.bold = True
            r_s.font.size = Pt(9.5)
            r_snip = p_s.add_run(f"\"{src['snippet'][:220]}...\"")
            r_snip.font.italic = True
            r_snip.font.size = Pt(8.5)
            r_snip.font.color.rgb = RGBColor(71, 85, 105)

    # Section 5: Engineering Recommendations
    h5 = doc.add_heading("5. Actionable Engineering Recommendations", level=1)
    h5.style.font.color.rgb = RGBColor(15, 23, 42)
    for rec in data.get("recommendations", []):
        p_r = doc.add_paragraph(style='List Number')
        r_r = p_r.add_run(rec)
        r_r.font.size = Pt(10)

    # Sign-off block
    doc.add_paragraph()
    doc.add_paragraph()
    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.rows[0].cells[0].text = "_____________________________________\nLead Inspection Engineer"
    sign_table.rows[0].cells[1].text = "_____________________________________\nChief Plant Superintendent / Authority"
    sign_table.rows[1].cells[0].text = f"Date: {datetime.utcnow().strftime('%Y-%m-%d')}"
    sign_table.rows[1].cells[1].text = "Status: ACTION REQUIRED / APPROVED WITH CONDITIONS"

    doc.save(str(output_path))
    return str(output_path)

def generate_pdf_approval_report(data: Dict[str, Any], output_path: Path) -> str:
    """
    Generates a formal, publication-grade PDF inspection report using ReportLab.
    """
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0F172A'),
        fontName='Helvetica-Bold'
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        fontName='Helvetica-Bold',
        spaceBefore=10,
        spaceAfter=4
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        fontName='Helvetica'
    )
    meta_style = ParagraphStyle(
        'MetaText',
        parent=styles['Normal'],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#475569'),
        fontName='Helvetica'
    )

    elements = []

    # Enclave classification banner
    elements.append(Paragraph("<b>SOVEREIGN INDUSTRIAL & DEFENCE ENCLAVE | AIR-GAPPED WORKBENCH</b>", meta_style))
    elements.append(Spacer(1, 4))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceBefore=2, spaceAfter=8))

    # Title
    elements.append(Paragraph(f"TECHNICAL EVALUATION REPORT", title_style))
    elements.append(Paragraph(f"<b>Workflow:</b> {data.get('task_type', 'Document Analysis')} | <b>Directive:</b> {data.get('requirement', '')[:100]}", meta_style))
    elements.append(Spacer(1, 8))

    # Metadata Table
    meta_table_data = [
        [Paragraph("<b>Clearance:</b>", meta_style), Paragraph("CONFIDENTIAL / RESTRICTED", meta_style),
         Paragraph("<b>Date:</b>", meta_style), Paragraph(datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"), meta_style)],
        [Paragraph("<b>Input Files:</b>", meta_style), Paragraph(", ".join(data.get("input_files", [])) or "Operational Data", meta_style),
         Paragraph("<b>Enclave ID:</b>", meta_style), Paragraph("SOV-LOCAL-AIRGAP-NODE-01", meta_style)]
    ]
    t_meta = Table(meta_table_data, colWidths=[65, 200, 65, 190])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    elements.append(t_meta)
    elements.append(Spacer(1, 12))

    # Executive Summary
    elements.append(Paragraph("1. Executive Summary", h2_style))
    elements.append(Paragraph(data.get("executive_summary", "Evaluation completed."), body_style))
    elements.append(Spacer(1, 10))

    # Key Findings
    elements.append(Paragraph("2. Technical Findings", h2_style))
    for f in data.get("key_findings", []):
        elements.append(Paragraph(f"• {f}", body_style))
    elements.append(Spacer(1, 10))

    # Critical Issues Table
    critical_issues = data.get("critical_issues", [])
    if critical_issues:
        elements.append(Paragraph("3. Safety Deviations & Non-Conformances", h2_style))
        table_rows = [["Severity", "Component", "Observed Value / Deviation", "Required Threshold"]]
        for ci in critical_issues:
            table_rows.append([
                ci.get("severity", "WARNING"),
                ci.get("component", ci.get("column", "System")),
                ci.get("deviation", ci.get("message", "")),
                ci.get("risk", ci.get("expected_range", "N/A"))
            ])

        t_issues = Table(table_rows, colWidths=[65, 110, 215, 130])
        t_issues.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 8),
            ('FONTSIZE', (0,1), (-1,-1), 8),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ]))
        elements.append(t_issues)
        elements.append(Spacer(1, 10))

    # Sources Used (Qdrant RAG)
    sources = data.get("sources_used", [])
    if sources:
        elements.append(Paragraph("4. Private Knowledge Base Citations (Qdrant RAG)", h2_style))
        for src in sources:
            elements.append(Paragraph(f"<b>[Source: {src['document_name']} - {src.get('section', '')}]</b> (Relevance: {int(src['relevance_score']*100)}%)", meta_style))
            elements.append(Paragraph(f"<i>\"{src['snippet'][:180]}...\"</i>", body_style))
            elements.append(Spacer(1, 4))

    # Recommendations
    elements.append(Paragraph("5. Operational Recommendations", h2_style))
    for idx, rec in enumerate(data.get("recommendations", []), 1):
        elements.append(Paragraph(f"<b>{idx}.</b> {rec}", body_style))

    doc.build(elements)
    return str(output_path)
